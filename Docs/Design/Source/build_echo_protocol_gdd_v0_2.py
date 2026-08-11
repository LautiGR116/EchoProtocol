#!/usr/bin/env python3
"""Build the editable Echo Protocol GDD v0.2 DOCX.

Design system:
- Base preset: compact_reference_guide.
- Named override: Echo Protocol clinical-memory editorial dossier.
- Named override: Arial Narrow/Arial typography with a reproducible doorway mark.
- Named override: compact table text for dense production matrices.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "Docs/Design/Echo Protocol - Game Design Document v0.2.docx"

# Base preset: compact_reference_guide.
PAGE_W_DXA = 12240
PAGE_H_DXA = 15840
MARGIN_DXA = 1440
CONTENT_W_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}

# Echo Protocol named palette override, sampled from the v0.2 direction board:
# institutional green-grey, aged domestic warmth, restrained signal cyan.
INK = "202726"
CHARCOAL = "171D1D"
MIDNIGHT = "090E0F"
CYAN = "78B9BD"
CYAN_DARK = "456F72"
ICE = "DDE4DF"
FROST = "ECEDE7"
STEEL = "65716E"
LINE = "ADB4AA"
MEMORY = "A8784F"
MEMORY_LIGHT = "E9DFD0"
ALERT = "763C40"
ALERT_LIGHT = "E6D8D4"
WHITE = "FFFFFF"
BLACK = "000000"

FONT_BODY = "Arial"
FONT_DISPLAY = "Arial Narrow"
FONT_MONO = "Courier New"
FONT_MEMORY = "Georgia"

BRAND_MARK = ROOT / "Docs/Design/Brand/EchoProtocol_Mark.png"


def _ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_width(parent, tag: str, width_dxa: int) -> None:
    width = _ensure_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(int(width_dxa)))


def _set_cell_margins(cell, margins_dxa: dict[str, int]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = _ensure_child(tc_pr, "w:tcMar")
    for side in ("top", "bottom", "start", "end"):
        margin = _ensure_child(tc_mar, f"w:{side}")
        margin.set(qn("w:w"), str(int(margins_dxa[side])))
        margin.set(qn("w:type"), "dxa")


def apply_table_geometry(
    table,
    column_widths_dxa,
    *,
    table_width_dxa=None,
    indent_dxa=None,
    cell_margins_dxa=None,
) -> None:
    """Apply deterministic Word widths, grid columns, indents and cell margins."""

    widths = [int(width) for width in column_widths_dxa]
    width_total = int(table_width_dxa if table_width_dxa is not None else sum(widths))
    if not widths or any(width <= 0 for width in widths) or sum(widths) != width_total:
        raise ValueError("invalid table geometry")

    margins = dict(CELL_MARGINS)
    if cell_margins_dxa:
        margins.update({key: int(value) for key, value in cell_margins_dxa.items()})
    resolved_indent = margins["start"] if indent_dxa is None else int(indent_dxa)

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    _set_width(tbl_pr, "w:tblW", width_total)

    table_indent = _ensure_child(tbl_pr, "w:tblInd")
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(resolved_indent))
    _ensure_child(tbl_pr, "w:tblLayout").set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)
    for row in table.rows:
        if len(row.cells) != len(widths):
            raise ValueError("table rows must remain unmerged while geometry is applied")
        row.height = None
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            _set_width(cell._tc.get_or_add_tcPr(), "w:tcW", width)
            _set_cell_margins(cell, margins)


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, *, name=FONT_BODY, size=None, color=INK, bold=None, italic=None, caps=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if caps is not None:
        run.font.all_caps = caps
    return run


def set_style_font(style, *, name=FONT_BODY, size=11, color=INK, bold=False, italic=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold
    style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def set_paragraph_shading(paragraph, fill: str):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, *, side="left", color=CYAN, size=18, space=8):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    edge = pbdr.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        pbdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge_name, spec in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), spec.get("val", "single"))
        edge.set(qn("w:sz"), str(spec.get("sz", 6)))
        edge.set(qn("w:space"), str(spec.get("space", 0)))
        edge.set(qn("w:color"), spec.get("color", LINE))


def set_table_borders(table, *, color=LINE, size=5, inside=True):
    edges = {
        "top": {"color": color, "sz": size},
        "left": {"color": color, "sz": size},
        "bottom": {"color": color, "sz": size},
        "right": {"color": color, "sz": size},
    }
    if inside:
        edges["insideH"] = {"color": color, "sz": size}
        edges["insideV"] = {"color": color, "sz": size}
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge_name, spec in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(spec["sz"]))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), spec["color"])


def set_row_repeat(row):
    trpr = row._tr.get_or_add_trPr()
    header = trpr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        trpr.append(header)
    header.set(qn("w:val"), "true")


def set_row_cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = trpr.find(qn("w:cantSplit"))
    if el is None:
        el = OxmlElement("w:cantSplit")
        trpr.append(el)


def remove_table_cell_paragraph(cell):
    p = cell.paragraphs[0]
    for run in list(p.runs):
        p._p.remove(run._r)
    return p


def style_cell_text(cell, *, color=INK, size=9.2, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font=FONT_BODY):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.12
        for run in p.runs:
            set_run_font(run, name=font, size=size, color=color, bold=bold)


def add_hyperlink(paragraph, text: str, url: str, *, color=CYAN_DARK, underline=True, size=8.5):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT_BODY)
    rpr.append(rfonts)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    rpr.append(color_el)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    new_run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_picture_with_alt(run, path: Path, *, width, title: str, description: str):
    shape = run.add_picture(str(path), width=width)
    shape._inline.docPr.set("title", title)
    shape._inline.docPr.set("descr", description)
    return shape


def add_page_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), STEEL)
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def add_page_count_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "NUMPAGES")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), STEEL)
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "18"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    set_style_font(normal, name=FONT_BODY, size=10.4, color=INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.keep_together = False
    normal.paragraph_format.widow_control = True

    h1 = styles["Heading 1"]
    set_style_font(h1, name=FONT_DISPLAY, size=20, color=CHARCOAL, bold=True)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.line_spacing = 1.0
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True

    h2 = styles["Heading 2"]
    set_style_font(h2, name=FONT_DISPLAY, size=15, color=CYAN_DARK, bold=True)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(7)
    h2.paragraph_format.line_spacing = 1.0
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True
    h2_ppr = h2.element.get_or_add_pPr()
    h2_borders = OxmlElement("w:pBdr")
    h2_left = OxmlElement("w:left")
    h2_left.set(qn("w:val"), "single")
    h2_left.set(qn("w:sz"), "12")
    h2_left.set(qn("w:space"), "7")
    h2_left.set(qn("w:color"), MEMORY)
    h2_borders.append(h2_left)
    h2_ppr.append(h2_borders)

    h3 = styles["Heading 3"]
    set_style_font(h3, name=FONT_DISPLAY, size=12.5, color=STEEL, bold=True)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(5)
    h3.paragraph_format.line_spacing = 1.0
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.keep_together = True

    custom_styles = [
        ("EP Kicker", FONT_MONO, 7.8, CYAN_DARK, True, False),
        ("EP Deck", FONT_BODY, 12, STEEL, False, False),
        ("EP Caption", FONT_BODY, 8.5, STEEL, False, True),
        ("EP Table", FONT_BODY, 9.2, INK, False, False),
        ("EP Table Header", FONT_DISPLAY, 8.8, WHITE, True, False),
        ("EP Mono", FONT_MONO, 9.2, CYAN_DARK, False, False),
        ("EP Cover Kicker", FONT_MONO, 7.8, CYAN, True, False),
        ("EP Cover Title", FONT_DISPLAY, 42, WHITE, True, False),
        ("EP Cover Subtitle", FONT_BODY, 11, ICE, False, False),
        ("EP Part Label", FONT_MONO, 7.8, MEMORY, True, False),
        ("EP Small", FONT_BODY, 8.5, STEEL, False, False),
        ("EP Memory", FONT_MEMORY, 10.5, ALERT, False, True),
    ]
    for name, font, size, color, bold, italic in custom_styles:
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        set_style_font(style, name=font, size=size, color=color, bold=bold, italic=italic)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.0 if "Cover" in name or name in ("EP Kicker", "EP Part Label") else 1.15


def add_custom_numbering(doc: Document):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs, default=0) + 1
    next_num = max(existing_num, default=0) + 1

    def make_definition(abs_id: int, num_id: int, fmt: str, level_text: str):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        numfmt = OxmlElement("w:numFmt")
        numfmt.set(qn("w:val"), fmt)
        lvl.append(numfmt)
        lvltext = OxmlElement("w:lvlText")
        lvltext.set(qn("w:val"), level_text)
        lvl.append(lvltext)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "271")
        ppr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        lvl.append(ppr)
        if fmt == "bullet":
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), FONT_BODY)
            rfonts.set(qn("w:hAnsi"), FONT_BODY)
            rpr.append(rfonts)
            lvl.append(rpr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        aid = OxmlElement("w:abstractNumId")
        aid.set(qn("w:val"), str(abs_id))
        num.append(aid)
        numbering.append(num)

    make_definition(next_abs, next_num, "bullet", "•")
    make_definition(next_abs + 1, next_num + 1, "decimal", "%1.")
    return next_num, next_num + 1


def apply_numbering(paragraph, num_id: int, ilvl=0):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    numid_el = OxmlElement("w:numId")
    numid_el.set(qn("w:val"), str(num_id))
    numpr.append(ilvl_el)
    numpr.append(numid_el)
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def add_bullet(doc, text: str, bullet_num_id: int, *, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    apply_numbering(p, bullet_num_id)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_number(doc, text: str, number_num_id: int):
    p = doc.add_paragraph()
    apply_numbering(p, number_num_id)
    set_run_font(p.add_run(text))
    return p


def add_kicker(doc, text: str, *, color=None):
    p = doc.add_paragraph(style="EP Part Label")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    run = p.add_run(text.upper())
    set_run_font(run, size=8.5, color=color or MEMORY, bold=True, caps=True)
    return p


def add_body(doc, text: str, *, bold_lead: str | None = None, italic=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]), italic=italic)
    else:
        set_run_font(p.add_run(text), italic=italic)
    return p


def add_small(doc, text: str, *, italic=False):
    p = doc.add_paragraph(style="EP Small")
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(text), size=8.5, color=STEEL, italic=italic)
    return p


def add_callout(doc, label: str, text: str, *, fill=ICE, accent=CYAN_DARK, text_color=INK, trailing_space=True):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    no = {"val": "nil", "sz": 0, "color": fill}
    set_cell_border(cell, top=no, bottom=no, right=no, left={"val": "single", "sz": 24, "color": accent})
    p = remove_table_cell_paragraph(cell)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(label.upper() + "  "), size=8.5, color=accent, bold=True, caps=True)
    set_run_font(p.add_run(text), size=10.5, color=text_color, bold=True if len(text) < 150 else False)
    apply_table_geometry(table, [CONTENT_W_DXA], table_width_dxa=CONTENT_W_DXA, indent_dxa=180, cell_margins_dxa={"top": 110, "bottom": 110, "start": 180, "end": 180})
    if trailing_space:
        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, widths, *, header_fill=CHARCOAL, zebra=True, font_size=9.2, compact=False):
    cols = len(headers)
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(value)
        set_cell_shading(cell, header_fill)
        style_cell_text(cell, color=WHITE, size=8.7, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font=FONT_DISPLAY)
    set_row_repeat(table.rows[0])
    set_row_cant_split(table.rows[0])
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        for c_idx, value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(value)
            if zebra and r_idx % 2 == 1:
                set_cell_shading(cell, FROST)
            style_cell_text(cell, size=font_size)
    apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=TABLE_INDENT_DXA, cell_margins_dxa=CELL_MARGINS)
    set_table_borders(table, color=LINE, size=4, inside=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2 if compact else 5)
    return table


def add_label_value_table(doc, rows, *, label_w=2520, value_w=6840, font_size=9.4):
    table = doc.add_table(rows=1, cols=2)
    for idx, (label, value) in enumerate(rows):
        row = table.rows[0] if idx == 0 else table.add_row()
        set_row_cant_split(row)
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], ICE)
        style_cell_text(row.cells[0], color=CYAN_DARK, size=8.7, bold=True, font=FONT_DISPLAY)
        style_cell_text(row.cells[1], size=font_size)
    apply_table_geometry(table, [label_w, value_w], table_width_dxa=label_w + value_w, indent_dxa=TABLE_INDENT_DXA, cell_margins_dxa=CELL_MARGINS)
    set_table_borders(table, color=LINE, size=4, inside=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    return table


def add_flow_strip(doc, steps, *, fill=ICE, accent=CYAN_DARK):
    widths = [CONTENT_W_DXA // len(steps)] * len(steps)
    widths[-1] += CONTENT_W_DXA - sum(widths)
    table = doc.add_table(rows=1, cols=len(steps))
    for i, step in enumerate(steps):
        cell = table.cell(0, i)
        set_cell_shading(cell, fill if i % 2 == 0 else FROST)
        p = remove_table_cell_paragraph(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(f"{i + 1:02d}\n"), size=8, color=accent, bold=True)
        set_run_font(p.add_run(step.upper()), size=8.7, color=INK, bold=True)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell, top={"color": LINE, "sz": 4}, bottom={"color": LINE, "sz": 4}, left={"color": LINE, "sz": 4}, right={"color": LINE, "sz": 4})
    apply_table_geometry(table, widths, table_width_dxa=CONTENT_W_DXA, indent_dxa=70, cell_margins_dxa={"top": 90, "bottom": 90, "start": 70, "end": 70})
    p = doc.add_paragraph(style="EP Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("El ciclo vuelve a comenzar con mayor conocimiento y menor confianza."), size=8.5, color=STEEL, italic=True)
    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.paragraph_format.space_after = Pt(0)


def add_part_opening(doc, number: str, title: str, subtitle: str, *, page_break_before=False):
    anchor = doc.add_paragraph()
    anchor.paragraph_format.page_break_before = page_break_before
    anchor.paragraph_format.space_before = Pt(0)
    anchor.paragraph_format.space_after = Pt(3)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_row_cant_split(table.rows[0])
    left, right = table.rows[0].cells
    set_cell_shading(left, MIDNIGHT)
    set_cell_shading(right, MEMORY_LIGHT)

    lp = remove_table_cell_paragraph(left)
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.space_before = Pt(14)
    lp.paragraph_format.space_after = Pt(14)
    set_run_font(lp.add_run(number), name=FONT_DISPLAY, size=32, color=CYAN, bold=True)

    rp = remove_table_cell_paragraph(right)
    rp.paragraph_format.space_before = Pt(7)
    rp.paragraph_format.space_after = Pt(2)
    set_run_font(rp.add_run(title.upper()), name=FONT_DISPLAY, size=20, color=CHARCOAL, bold=True, caps=True)
    deck = right.add_paragraph()
    deck.paragraph_format.space_before = Pt(0)
    deck.paragraph_format.space_after = Pt(7)
    set_run_font(deck.add_run(subtitle), name=FONT_BODY, size=9.8, color=STEEL)

    apply_table_geometry(
        table,
        [1320, 8040],
        table_width_dxa=CONTENT_W_DXA,
        indent_dxa=150,
        cell_margins_dxa={"top": 50, "bottom": 50, "start": 150, "end": 150},
    )
    set_table_borders(table, color=MEMORY, size=5, inside=True)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def add_status_line(doc, label: str, value: str, *, color=CYAN_DARK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    set_paragraph_shading(p, FROST)
    set_run_font(p.add_run(f"  {label.upper()}  "), size=8.2, color=WHITE, bold=True)
    set_run_font(p.runs[0], size=8.2, color=WHITE, bold=True)
    p.runs[0]._element.get_or_add_rPr()
    # Keep the chip as text; the paragraph shading provides a restrained band.
    set_run_font(p.add_run(value), size=8.5, color=color, bold=True)
    return p


def cover_page(doc):
    p = doc.add_paragraph(style="EP Cover Kicker")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    set_run_font(
        p.add_run("UNIVERSIDAD AUSTRAL  /  LABORATORIO DE DESARROLLO DE VIDEOJUEGOS"),
        name=FONT_MONO,
        size=7.5,
        color=MEMORY,
        bold=True,
        caps=True,
    )

    mark = doc.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mark.paragraph_format.space_before = Pt(3)
    mark.paragraph_format.space_after = Pt(0)
    set_paragraph_shading(mark, MIDNIGHT)
    set_paragraph_border(mark, side="left", color=MEMORY, size=16, space=6)
    add_picture_with_alt(
        mark.add_run(),
        BRAND_MARK,
        width=Inches(0.78),
        title="Echo Protocol mark",
        description="Tres marcos de puerta concéntricos rodean una figura humana.",
    )

    title = doc.add_paragraph(style="EP Cover Title")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    title.paragraph_format.line_spacing = 0.82
    set_paragraph_shading(title, MIDNIGHT)
    set_run_font(title.add_run("ECHO\nPROTOCOL"), name=FONT_DISPLAY, size=47, color=WHITE, bold=True, caps=True)

    signal = doc.add_paragraph()
    signal.paragraph_format.space_before = Pt(0)
    signal.paragraph_format.space_after = Pt(0)
    signal.paragraph_format.line_spacing = Pt(12)
    set_paragraph_shading(signal, MIDNIGHT)
    set_run_font(signal.add_run("RECALL  /  ECHO  /  MEMORY  /  DENIAL"), name=FONT_MONO, size=7.4, color=CYAN, bold=True)

    subtitle = doc.add_paragraph(style="EP Cover Subtitle")
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(0)
    subtitle.paragraph_format.line_spacing = Pt(14)
    set_paragraph_shading(subtitle, MIDNIGHT)
    set_run_font(subtitle.add_run("GAME DESIGN DOCUMENT  0.2\nTerror psicológico y cooperación temporal en primera persona"), size=10.8, color=ICE)

    protocol = doc.add_table(rows=1, cols=4)
    protocol.alignment = WD_TABLE_ALIGNMENT.LEFT
    stages = [
        ("01", "OBSERVAR"),
        ("02", "RECORDAR"),
        ("03", "REPETIR"),
        ("04", "NEGAR"),
    ]
    for index, (code, label) in enumerate(stages):
        cell = protocol.cell(0, index)
        set_cell_shading(cell, ICE if index % 2 == 0 else FROST)
        pc = remove_table_cell_paragraph(cell)
        pc.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pc.paragraph_format.space_before = Pt(7)
        pc.paragraph_format.space_after = Pt(7)
        set_run_font(pc.add_run(code + "\n"), name=FONT_MONO, size=7, color=MEMORY, bold=True)
        set_run_font(pc.add_run(label), name=FONT_DISPLAY, size=9.5, color=INK, bold=True)
    apply_table_geometry(protocol, [2340] * 4, table_width_dxa=CONTENT_W_DXA, indent_dxa=120, cell_margins_dxa={"top": 60, "bottom": 60, "start": 120, "end": 120})
    set_table_borders(protocol, color=LINE, size=4, inside=True)

    divider = doc.add_paragraph()
    divider.paragraph_format.space_before = Pt(7)
    divider.paragraph_format.space_after = Pt(7)
    set_paragraph_border(divider, side="bottom", color=MEMORY, size=12, space=4)

    meta = doc.add_table(rows=2, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    values = [
        ("DOCUMENTO", "EP-GDD-001  /  VERSIÓN 0.2"),
        ("ESTADO", "DISEÑO + PROTOTIPO"),
        ("AUTOR", "LAUTARO REINOSO"),
        ("FECHA", "11.08.2026  /  2.º CUATRIMESTRE"),
    ]
    for idx, (label, value) in enumerate(values):
        cell = meta.cell(idx // 2, idx % 2)
        set_cell_shading(cell, MEMORY_LIGHT if idx < 2 else FROST)
        pcell = remove_table_cell_paragraph(cell)
        pcell.paragraph_format.space_before = Pt(2)
        pcell.paragraph_format.space_after = Pt(2)
        set_run_font(pcell.add_run(label + "\n"), name=FONT_MONO, size=6.8, color=ALERT, bold=True)
        set_run_font(pcell.add_run(value), name=FONT_BODY, size=8.5, color=INK, bold=True)
    apply_table_geometry(meta, [4680, 4680], table_width_dxa=CONTENT_W_DXA, indent_dxa=120, cell_margins_dxa={"top": 75, "bottom": 75, "start": 120, "end": 120})
    set_table_borders(meta, color=LINE, size=4, inside=True)

    p = doc.add_paragraph(style="EP Memory")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(0)
    set_paragraph_border(p, side="right", color=ALERT, size=14, space=6)
    set_run_font(
        p.add_run("Los Echoes recuerdan correctamente. El protagonista no."),
        name=FONT_MEMORY,
        size=10.5,
        color=ALERT,
        bold=True,
        italic=True,
    )


def document_control_page(doc, bullet_id):
    add_part_opening(doc, "00", "Control del documento", "Versión, alcance y mapa de lectura.")
    add_table(
        doc,
        ["Versión", "Fecha", "Estado", "Responsable"],
        [
            ["0.1", "08/08/2026", "Visión inicial", "Lautaro Reinoso"],
            ["0.2", "11/08/2026", "Diseño actualizado", "Lautaro Reinoso"],
        ],
        [1200, 1800, 2460, 3900],
        header_fill=CHARCOAL,
        zebra=False,
        font_size=9.5,
    )
    add_callout(
        doc,
        "Documento vivo",
        "Esta versión sincroniza la visión con el prototipo y sus playtests. El repositorio conserva la autoridad sobre lo implementado; EN PRUEBA y TBD siguen requiriendo evidencia.",
        fill=ICE,
        accent=CYAN_DARK,
    )
    doc.add_heading("Estructura basada en Unity", level=2)
    add_table(
        doc,
        ["Bloque", "Contenido"],
        [
            ("01  Introducción", "Identidad, concepto, público, experiencia, features y pilares."),
            ("02  Gameplay", "Jugador, controles, core loop, RECALL, Echoes, puzzles, flujo, horror, narrativa y UI."),
            ("03  Arte y audio", "Dirección audiovisual, progresión, paleta y tablero de intención."),
            ("04  Producción", "Especificaciones, alcance, estado real, próximos gates, riesgos y decisiones abiertas."),
        ],
        [2200, 7160],
        header_fill=CYAN_DARK,
        font_size=9.2,
    )
    doc.add_heading("Convenciones", level=2)
    add_bullet(doc, "IMPLEMENTADO: existe en el repositorio y fue inspeccionado.", bullet_id, bold_prefix="IMPLEMENTADO:")
    add_bullet(doc, "VALIDADO: superó evidencia objetiva o un playtest humano nombrado.", bullet_id, bold_prefix="VALIDADO:")
    add_bullet(doc, "DEFINIDO: dirección aprobada, aunque todavía no esté implementada.", bullet_id, bold_prefix="DEFINIDO:")
    add_bullet(doc, "EN PRUEBA: hipótesis que requiere playtesting.", bullet_id, bold_prefix="EN PRUEBA:")
    add_bullet(doc, "TBD: detalle deliberadamente abierto; no bloquea el prototipo.", bullet_id, bold_prefix="TBD:")
    add_small(doc, "La v0.1 se conserva sin cambios. Esta v0.2 resuelve contradicciones del prototipo y reemplaza el cronograma provisional por estado y gates reales.", italic=True)


def intro_page_one(doc):
    add_part_opening(doc, "01", "Introducción", "Qué es Echo Protocol, para quién se diseña y qué promete.")
    doc.add_heading("1.1 Ficha del proyecto", level=2)
    add_label_value_table(
        doc,
        [
            ("Título provisional", "Echo Protocol"),
            ("Autor / equipo", "Lautaro Reinoso / 1 desarrollador"),
            ("Género", "Psychological Horror + Temporal Puzzle + Narrative Adventure"),
            ("Modo / jugadores", "Single-player / 1 jugador"),
            ("Formato y cámara", "3D / primera persona"),
            ("Plataformas", "PC: Windows y macOS"),
            ("Motor", "Unity 6.3 LTS + Universal Render Pipeline (URP)"),
            ("Input inicial", "Keyboard + Mouse"),
            ("Duración objetivo", "30-45 minutos; reducible antes que sacrificar calidad"),
            ("Estado", "Versión 0.2 / Prototipo M0-M6 + dirección SD1"),
        ],
        label_w=2400,
        value_w=6960,
        font_size=9.2,
    )
    doc.add_heading("1.2 Concepto", level=2)
    add_body(doc, "El protagonista despierta en una instalación desconocida y extrañamente familiar. Para avanzar debe superar pruebas con RECALL: el sistema registra cada intento y, al reiniciar el ciclo, crea un Echo que reproduce sus acciones pasadas.")
    add_body(doc, "Los Echoes son herramientas precisas y confiables. Más adelante pueden revelar ciclos verdaderos que el protagonista no recuerda. La instalación se mezcla con espacios, sonidos y figuras de su vida hasta demostrar que el error no está en las grabaciones, sino en su memoria y percepción.")
    add_callout(doc, "High concept", "Versiones pasadas cooperan con el jugador mediante reproducciones exactas; cuando una grabación muestra algo que el protagonista no recuerda, el sistema se convierte en testigo de su negación.", fill=CHARCOAL, accent=CYAN, text_color=WHITE)


def intro_page_two(doc, bullet_id):
    add_kicker(doc, "01 / Introducción")
    doc.add_heading("1.3 Logline, promesa y USP", level=2)
    add_body(doc, "Un juego breve de terror psicológico donde la cooperación temporal se convierte en una investigación sobre culpa, recuerdo y negación.")
    add_callout(doc, "North star", "Era el juego donde tus versiones pasadas recordaban correctamente lo que vos no podías recordar.", fill=ICE, accent=CYAN_DARK)
    add_body(doc, "La propuesta diferencial no es sumar un monstruo a un puzzle game. Es usar un testigo temporal exacto dentro de una experiencia de horror que ataca la percepción del jugador mediante espacio, sonido, memoria y evidencia parcial.")

    doc.add_heading("1.4 Público objetivo", level=2)
    add_body(doc, "Jugadores adolescentes mayores y adultos interesados en terror psicológico, misterios narrativos, puzzles accesibles y experiencias breves de alta intensidad. No apunta a quienes buscan combate, progresión RPG o gore como atractivo central.")
    add_bullet(doc, "Temas sensibles: muerte, culpa, infancia, accidentes, pérdida, duelo y memoria distorsionada.", bullet_id)
    add_bullet(doc, "Contenido fuerte por implicación; representación explícita aproximada 3-4/5.", bullet_id)
    add_bullet(doc, "Audio principal en inglés con subtítulos en español e inglés.", bullet_id)

    doc.add_heading("1.5 Experiencia del jugador", level=2)
    add_flow_strip(doc, ["Curiosidad", "Inquietud", "Duda", "Paranoia", "Terror", "Aceptación"], fill=ICE, accent=CYAN_DARK)
    add_table(
        doc,
        ["Momento", "Pensamiento buscado"],
        [
            ("Inicio", "Qué interesante esta mecánica."),
            ("Primer cambio", "¿Eso vino del juego o de mi casa?"),
            ("Contradicción", "El Echo hizo exactamente eso. ¿Por qué no lo recuerdo?"),
            ("Revelación", "Si el Echo tiene razón, ¿qué pasó realmente?"),
        ],
        [2200, 7160],
        header_fill=CYAN_DARK,
        font_size=9.3,
    )


def intro_page_three(doc):
    add_kicker(doc, "01 / Introducción")
    doc.add_heading("1.6 Pilares de diseño", level=2)
    add_table(
        doc,
        ["Pilar", "Directriz operativa"],
        [
            ("Trust Before Fear", "La regla debe volverse seguridad antes de que el contexto haga dudar; el Echo no se corrompe."),
            ("The Mechanic Is the Horror", "RECALL y el Echo exacto exponen lo que el protagonista niega o no recuerda."),
            ("Player-facing Horror", "El miedo debe alterar la percepción y conducta del jugador real; el lore de un objeto no alcanza."),
            ("Strong Implication, Limited Explicitness", "El jugador reconstruye la historia a través de objetos, sonidos, espacios, motivos y asociaciones."),
            ("Small Scope, High Polish", "Un juego corto, legible y pulido tiene prioridad sobre más contenido incompleto."),
            ("Sound as the Main Weapon", "Dirección, proximidad, motivos, música, contraste y repetición hacen dudar del entorno real sin exigir volumen alto."),
        ],
        [2600, 6760],
        header_fill=CHARCOAL,
        font_size=9.2,
    )

    doc.add_heading("1.7 Características clave", level=2)
    add_table(
        doc,
        ["Feature", "Valor para el jugador", "Estado"],
        [
            ("RECALL + Echo exacto", "Convierte cada intento iniciado en material de cooperación temporal.", "VALIDADO"),
            ("Un Echo", "Lenguaje actual confiable; ampliar sólo si un puzzle concreto lo justifica.", "VALIDADO"),
            ("Audio espacial", "Hace dudar entre espacio virtual, memoria y entorno físico.", "DEFINIDO"),
            ("Puzzles vulnerables", "Obligan a esperar, mirar fuera, regresar o usar RECALL bajo presión.", "DEFINIDO"),
            ("Hub alterable", "Recontextualiza espacios conocidos mediante backtracking breve.", "PLANIFICADO"),
            ("Objetos móviles por Echo", "Ampliaría puzzles, pero aumenta conflictos de timeline.", "NO MVP"),
        ],
        [2500, 4960, 1900],
        header_fill=CYAN_DARK,
        font_size=8.9,
    )
    doc.add_heading("1.8 ¿Por qué es atractivo?", level=2)
    add_body(doc, "El jugador obtiene satisfacción al planificar, ejecutar y observar cómo una versión pasada completa una parte del plan. El mismo dominio crea vulnerabilidad: cuanto mejor comprende la grabación, más inquietante resulta que el contexto o su propio recuerdo no encajen con ella.")
    add_callout(doc, "Alcance", "Objetivo: 30-45 minutos muy pulidos. Si el calendario exige recorte, 20 minutos excelentes tienen prioridad sobre 45 minutos mediocres.", fill=MEMORY_LIGHT, accent=MEMORY)


def gameplay_overview_page(doc, bullet_id):
    add_part_opening(doc, "02", "Gameplay", "Reglas, sistemas, flujo, tensión y narrativa jugable.")
    doc.add_heading("2.1 Objetivo del jugador", level=2)
    add_body(doc, "Atravesar una secuencia de pruebas, coordinar acciones con Echoes y reconstruir qué ocurrió en dos acontecimientos traumáticos conectados. El objetivo local es resolver cada cámara; el objetivo global es confrontar el recuerdo que mantiene activa la pesadilla.")

    doc.add_heading("2.2 Core loop", level=2)
    add_flow_strip(doc, ["Explorar", "Escuchar", "Entender", "Planificar", "Exponerse", "RECALL", "Cooperar", "Verificar"], fill=ICE, accent=CYAN_DARK)
    add_bullet(doc, "Cada ciclo aporta información espacial o temporal.", bullet_id)
    add_bullet(doc, "Cada Echo conserva la verdad mecánica mientras cambia su significado emocional.", bullet_id)
    add_bullet(doc, "El avance combina puzzle, exploración, audio, narrativa, consecuencia y recuperación.", bullet_id)

    doc.add_heading("2.3 Definiciones de diseño", level=2)
    add_table(
        doc,
        ["Término", "Definición"],
        [
            ("Ciclo", "Período desde que comienza el registro de una prueba hasta que se activa RECALL o se completa la secuencia."),
            ("RECALL", "Acción que restaura el estado base de la prueba y convierte el intento anterior en un Echo."),
            ("Echo", "Reproducción temporal fiel de los canales realmente registrados; nunca fabrica ni altera datos."),
            ("Discrepancia", "Cambio de contexto, percepción o evidencia que no corrompe la grabación."),
            ("Prueba", "Situación temporal con objetivo, vulnerabilidad sensorial y condición de salida; no necesariamente una sala de placas."),
        ],
        [1900, 7460],
        header_fill=CHARCOAL,
        font_size=9.0,
    )


def player_page(doc, bullet_id):
    add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.4 Definición del jugador", level=2)
    add_body(doc, "El protagonista conserva su identidad y conocimientos generales, pero presenta vacíos y reconstrucciones alrededor de determinados traumas. Su perspectiva es funcional para jugar, pero no completamente confiable para interpretar lo ocurrido.")
    add_body(doc, "No necesita nombre en la versión 0.2. El juego evita una biografía expositiva: la identidad se revela mediante Echoes, recuerdos, objetos, sonido y reacciones mínimas. El nombre se define sólo si mejora la voz o la claridad narrativa.")

    doc.add_heading("2.5 Propiedades y habilidades", level=2)
    add_table(
        doc,
        ["Implementado ahora", "Planificado o fuera del inicio"],
        [
            ("Caminar y mirar", "Combate y armas"),
            ("Interactuar", "Salto como mecánica"),
            ("Activar RECALL", "Recoger / soltar objetos livianos"),
            ("Colisión y gravedad", "Brazos/manos visibles"),
            ("Un Echo de movimiento", "HP, estadísticas o habilidades especiales"),
        ],
        [4680, 4680],
        header_fill=CYAN_DARK,
        font_size=9.1,
    )
    add_callout(doc, "Sprint", "No se incluye al inicio. Puede evaluarse como velocidad contextual o necesidad tardía durante playtesting, sin permitir que el jugador atraviese corriendo la tensión.", fill=FROST, accent=MEMORY)

    doc.add_heading("2.6 Controles provisionales", level=2)
    add_table(
        doc,
        ["Acción", "Control", "Feedback"],
        [
            ("Movimiento", "WASD", "Desplazamiento + pasos"),
            ("Mirar", "Mouse", "Cámara en primera persona"),
            ("Interactuar", "E", "Respuesta diegética; retículo final TBD"),
            ("RECALL", "R", "Cue audiovisual + restauración del ciclo"),
            ("Cursor", "Esc / click", "Liberar / recapturar en el prototipo"),
        ],
        [2600, 1800, 4960],
        header_fill=CHARCOAL,
        font_size=9.0,
    )
    add_small(doc, "Game feel: movimiento preciso y deliberado, interacción inmediata pero sobria, y un pulso de RECALL inequívoco sin romper la inmersión.", italic=True)


def recall_page(doc, number_id):
    add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.7 Sistema RECALL", level=2)
    add_body(doc, "Cada intento queda armado al restaurarse, pero la línea temporal comienza con el primer desplazamiento real y conserva un lead-in fijo de un segundo. Así no acumula espera accidental en el spawn y sí reproduce completas las pausas deliberadas posteriores. RECALL se activa con R o al límite actual de 60 segundos.")
    add_flow_strip(doc, ["Registrar", "Activar RECALL", "Restaurar", "Crear Echo", "Reproducir", "Continuar"], fill=ICE, accent=CYAN_DARK)

    doc.add_heading("Reglas funcionales", level=3)
    for item in [
        "Registrar actualmente posición y orientación horizontal a 20 muestras por segundo.",
        "Restaurar puertas, placas, mecanismos, props controlados y estado de la prueba a una base conocida.",
        "Reproducir el intento anterior con temporización suficientemente estable para resolver el puzzle.",
        "Mantener un Echo: sólo la última grabación completada. Ampliar únicamente con evidencia de puzzle.",
        "Permitir varios resets consecutivos sin degradar el estado del nivel.",
    ]:
        add_number(doc, item, number_id)

    doc.add_heading("Ejemplo mínimo", level=3)
    add_table(
        doc,
        ["Ciclo 1", "Ciclo 2"],
        [
            ("Jugador camina a la placa.", "Echo 1 camina a la placa."),
            ("Jugador permanece sobre la placa.", "Echo 1 mantiene la puerta abierta."),
            ("Jugador activa RECALL.", "Jugador atraviesa la puerta."),
        ],
        [4680, 4680],
        header_fill=CYAN_DARK,
        font_size=9.3,
    )
    add_callout(doc, "Potencial de horror", "Más adelante puede recuperarse otro ciclo verdadero que el protagonista no recuerda. Debe presentarse como una grabación distinta u omitida, nunca como una copia conocida que cambió arbitrariamente.", fill=ALERT_LIGHT, accent=ALERT)


def echo_page(doc, bullet_id):
    add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.8 Sistema de Echoes", level=2)
    add_body(doc, "Los Echoes son representaciones temporales del propio protagonista. Deben comunicar siempre 'ese soy yo'. No evolucionan hacia monstruos y no alteran datos capturados; la perturbación proviene de lo que una grabación verdadera revela sobre la memoria y su nuevo contexto.")
    add_table(
        doc,
        ["Regla confiable", "Horror compatible", "Efecto"],
        [
            ("Repite ruta y timing exactos", "El entorno cambia el significado de esa misma ruta", "La memoria queda en duda"),
            ("Termina y desaparece", "Pasos ajenos continúan después", "Presencia no confirmada"),
            ("Refleja canales implementados", "Un nuevo canal revela más verdad después de ser aprobado", "Testigo más completo"),
            ("Existe si hubo grabación", "Se recupera un ciclo omitido claramente distinto", "Vacío de memoria"),
            ("No mira por cuenta propia", "Una evidencia aparece donde ya estaba orientado", "Relectura inquietante"),
            ("Sonido funcional no miente", "Motivos de tiempo repiten una acción desde otro espacio", "Contradicción contextual"),
        ],
        [2750, 3310, 3300],
        header_fill=CHARCOAL,
        font_size=8.8,
    )

    doc.add_heading("Capacidades del MVP", level=3)
    add_bullet(doc, "Movimiento y orientación.", bullet_id)
    add_bullet(doc, "Ocupación física de pressure plates por Player y Echo.", bullet_id)
    add_bullet(doc, "Interacciones semánticas y terminales: planificadas, no implementadas.", bullet_id)
    add_bullet(doc, "Sin manipulación de objetos físicos móviles en el MVP.", bullet_id)
    add_body(doc, "La reproducción debe favorecer confiabilidad visual antes que simulación física perfecta. Cada canal futuro se agrega de forma explícita y verificable; lo no registrado no se atribuye al Echo.")
    add_callout(doc, "Regla crítica", "La fidelidad del Echo es absoluta para los datos capturados. El horror puede cambiar interpretación, contexto o evidencia; no la grabación.", fill=ICE, accent=CYAN_DARK)
    add_small(doc, "Los ciclos omitidos, acciones reprimidas y percepción distorsionada son verdad narrativa registrada, no corrupción aleatoria.", italic=True)


def puzzle_page(doc, bullet_id):
    add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.9 Diseño de puzzles", level=2)
    add_body(doc, "Dificultad objetivo: 3/5. Los puzzles requieren planificación sin bloquear el ritmo. Además de coordinar Echoes, cada situación debe crear una vulnerabilidad concreta: esperar, escuchar, mirar fuera, regresar, separarse del Echo o activar RECALL bajo presión.")
    add_table(
        doc,
        ["Etapa", "Configuración", "Aprendizaje"],
        [
            ("Validado", "Una placa, dos placas y secuencia A > B.", "Confianza y planificación temporal"),
            ("Vulnerabilidad", "Terminal que obliga a dar la espalda.", "Atención dividida"),
            ("Retorno", "Acción que exige cruzar un espacio conocido.", "Comparación y anticipación"),
            ("Separación", "El Echo trabaja fuera de vista.", "Dependencia sin confirmación visual"),
            ("Escucha", "Esperar o sostener una acción mientras algo se aproxima.", "Presión perceptual"),
            ("Memoria", "Una grabación exacta adquiere otro significado.", "Mecánica convertida en horror"),
        ],
        [1700, 4190, 3470],
        header_fill=CYAN_DARK,
        font_size=8.9,
    )
    doc.add_heading("Lenguaje de puzzle", level=3)
    add_body(doc, "Se recombinan pocos elementos, pero el número final de pruebas permanece TBD. No se rellenan 30-45 minutos con variantes de placas de dos minutos: puzzle, exposición, motivo audiovisual, consecuencia y recuperación se diseñan juntos.")

    doc.add_heading("Ayuda adaptativa", level=3)
    add_bullet(doc, "Una luz o sonido orienta la atención.", bullet_id)
    add_bullet(doc, "La voz formula una observación dentro de la ficción.", bullet_id)
    add_bullet(doc, "Un Echo puede revelar indirectamente una relación causal.", bullet_id)
    add_bullet(doc, "No existe un botón tradicional de 'Mostrar pista'.", bullet_id)

    doc.add_heading("Recompensas", level=3)
    add_body(doc, "No hay pickups ni economía. Las recompensas son acceso a una nueva cámara, comprensión de una regla, dominio de la coordinación y una revelación narrativa o perceptiva.")


def flow_level_ui_page(doc, bullet_id):
    add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.10 Flujo del juego y estructura de nivel", level=2)
    add_flow_strip(doc, ["Menú", "Inicio", "Hub", "Prueba", "RECALL", "Resolución", "Retorno", "Capítulo"], fill=FROST, accent=MEMORY)
    add_body(doc, "La estructura combina un hub pequeño con secuencias lineales. El hub aporta orientación, respiración, transición narrativa y backtracking limitado. Cada retorno puede alterar un detalle que el jugador creía conocer.")
    add_table(
        doc,
        ["Componente", "Definición"],
        [
            ("Inicio de aplicación", "New Game, Continue, Settings y Exit."),
            ("Modo", "Single-player narrativo; sin modos alternativos en v0.2."),
            ("Victoria local", "Resolver la prueba y atravesar su salida."),
            ("Fracaso", "Sin HP ni muerte convencional; el sueño puede forzar un nuevo ciclo."),
            ("Guardado", "Autosave entre habitaciones o transiciones relevantes."),
            ("Final", "Confrontación del recuerdo, fin de la pesadilla y corte a negro; no se muestra el despertar."),
        ],
        [2400, 6960],
        header_fill=CHARCOAL,
        font_size=9.0,
    )

    doc.add_heading("2.11 UI / UX", level=2)
    add_bullet(doc, "Retículo contextual discreto y feedback claro de interacción.", bullet_id)
    add_bullet(doc, "Indicador mínimo de RECALL; cantidad de Echoes sólo si mejora legibilidad.", bullet_id)
    add_bullet(doc, "Información preferentemente diegética: por ejemplo, RECOLLECTION 02 / READY en una pantalla del mundo.", bullet_id)
    add_bullet(doc, "Audio principal en inglés; subtítulos configurables en español e inglés.", bullet_id)
    add_bullet(doc, "Headphones / Speakers, calibración L/R, Ambience / SFX / Voice y Reduced Dynamics: requisitos planificados.", bullet_id)
    add_bullet(doc, "Copy de calibración: 'Designed for headphones. Set a comfortable volume; details do not require high volume.'", bullet_id)
    add_bullet(doc, "Sin barra de vida, minimap, quest tracker ni HUD permanente complejo.", bullet_id)


def horror_chapters_page(doc):
    add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.12 Diseño de horror y progresión", level=2)
    add_body(doc, "La experiencia establece seguridad mecánica y después presiona la percepción del jugador mediante sonido espacial, recurrencia, variantes, consecuencias y evidencia incompleta. Puede ser intensa sin confirmar una amenaza física ni corromper el Echo.")
    add_table(
        doc,
        ["Capítulo", "Emoción", "Gameplay / horror", "Evolución del espacio"],
        [
            ("I - Protocol", "Curiosidad", "Reglas confiables y primer impacto inquietante.", "Instalación clínica funcional"),
            ("II - Familiarity", "Inquietud", "Duda espacial, sonidos domésticos y consecuencias sutiles.", "Detalles inexplicablemente familiares"),
            ("III - Recollection", "Duda", "Presencia, hogar y tiempo invaden las pruebas; el Echo sigue exacto.", "Habitaciones y pasillos personales"),
            ("IV - Denial", "Paranoia", "Grabaciones verdaderas contradicen el relato del protagonista.", "Instalación y memoria coexisten"),
            ("V - Acceptance", "Terror / tristeza", "Convergencia de traumas, responsabilidad y desenlace.", "Arquitectura imposible pero reconocible"),
        ],
        [1700, 1400, 3510, 2750],
        header_fill=CYAN_DARK,
        font_size=8.35,
    )

    doc.add_heading("Economía del miedo", level=3)
    add_table(
        doc,
        ["Priorizar", "Limitar"],
        [
            ("Motivos recurrentes con variantes autorales", "Ruido constante o eventos aleatorios"),
            ("Dirección, proximidad, ataque y contraste", "Volumen alto como sustituto de diseño"),
            ("Chequear, no encontrar y hallar evidencia después", "Gore o persecuciones constantes"),
            ("Preparación > impacto > consecuencia > recuperación", "Explicaciones inmediatas de anomalías"),
        ],
        [4680, 4680],
        header_fill=CHARCOAL,
        font_size=9.0,
    )
    add_callout(doc, "Peligro físico", "Durante aproximadamente el 70-80 % del juego no existe peligro físico confirmado. Esto no implica baja intensidad: presencia, golpes fuertes construidos y duda trans-pantalla pueden sostener presión sin un enemigo sistémico.", fill=ALERT_LIGHT, accent=ALERT)


def narrative_page(doc, bullet_id):
    add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.13 Diseño narrativo", level=2)
    add_callout(doc, "Giro central", "Los Echoes reproducen correctamente lo sucedido. La memoria y percepción del protagonista son las que han sido alteradas.", fill=CHARCOAL, accent=CYAN, text_color=WHITE)
    add_body(doc, "La pesadilla conecta dos acontecimientos: durante la infancia, el protagonista estuvo involucrado en un accidente que terminó con la muerte de su hermano/a menor; en la adultez, otro accidente relacionado con su hijo recreó simbólicamente la misma decisión.")
    add_small(doc, "La instalación funciona como gramática de la pesadilla. Queda abierto si tuvo una contraparte real; el juego no debe confirmarlo mediante exposición directa.", italic=True)

    doc.add_heading("Patrón psicológico", level=3)
    add_flow_strip(doc, ["Provoca indirectamente", "Entra en pánico", "Huye", "Altera el recuerdo"], fill=MEMORY_LIGHT, accent=MEMORY)
    add_body(doc, "El juego no lo convierte simplemente en villano. Al final se entiende cerca del 80 % de los acontecimientos y su responsabilidad, pero puede seguir abierta la pregunta de cuánto podría haber evitado.")

    doc.add_heading("Voz, personajes y motivos", level=3)
    add_bullet(doc, "Una voz inicialmente profesional guía las pruebas: 'Continue to the next room.'", bullet_id)
    add_bullet(doc, "Al principio puede negar o contradecir hechos que el jugador acaba de observar.", bullet_id)
    add_bullet(doc, "SD1 prueba un susurro inteligible y head-locked: 'Can you hear me?' / '¿Me escuchás?'. No define todavía identidad ni relación narrativa.", bullet_id)
    add_bullet(doc, "Más adelante una frase íntima puede revelar que la voz pertenece a alguien de su vida; identidad específica TBD.", bullet_id)
    add_bullet(doc, "Figuras del pasado progresan de sombra a silueta, vidrio, reflejo, grabación y reconocimiento parcial.", bullet_id)
    add_bullet(doc, "Motivos recurrentes: cuento infantil, dibujo, melodía, objeto, frase, habitación y recuerdo.", bullet_id)
    add_bullet(doc, "No hay NPCs que expliquen la historia de forma convencional.", bullet_id)

    doc.add_heading("Final C - No Awakening Shown", level=3)
    add_body(doc, "El protagonista confronta el recuerdo, la pesadilla termina y la pantalla corta a negro. No se muestra si despierta. El cierre busca horror, tristeza y una incertidumbre pequeña, sin sumar otro gran giro que invalide lo anterior.")


def art_page(doc, bullet_id):
    add_part_opening(doc, "03", "Arte y visuales", "Una identidad clínica que se vuelve íntima, imposible y perturbadora.")
    doc.add_heading("3.1 Dirección visual", level=2)
    add_body(doc, "Estilo semi-realista: proporciones creíbles, materiales, iluminación y composición por encima del fotorealismo AAA. El deterioro no vuelve todo oscuro y rojo; vuelve el espacio cada vez más personal.")
    add_table(
        doc,
        ["Fase", "Color / material", "Señal narrativa"],
        [
            ("Protocol", "Blanco, gris, azul frío, vidrio y metal", "Control institucional"),
            ("Familiarity", "Pequeños tonos cálidos y materiales domésticos", "Recuerdo infiltrado"),
            ("Recollection", "Madera, textiles, fotografía y luz mixta", "Dos mundos coexistiendo"),
            ("Denial / Acceptance", "Paleta contaminada, sombras y arquitectura imposible", "Memoria ya inseparable del espacio"),
        ],
        [1900, 3460, 4000],
        header_fill=CYAN_DARK,
        font_size=8.9,
    )

    doc.add_heading("3.2 Echoes y cuerpo", level=2)
    add_bullet(doc, "Silueta humana completa y reconocible, con transparencia leve, desaturación y artefactos temporales contenidos.", bullet_id)
    add_bullet(doc, "Sin deformación monstruosa progresiva; el comportamiento lleva el horror.", bullet_id)
    add_bullet(doc, "Jugador en primera persona con brazos/manos; rostro del protagonista reservado.", bullet_id)

    doc.add_heading("3.3 Paleta unificadora", level=2)
    table = doc.add_table(rows=2, cols=5)
    swatches = [
        ("Institución", ICE),
        ("Sombra", CHARCOAL),
        ("Echo", CYAN),
        ("Hogar", MEMORY),
        ("Herida", ALERT),
    ]
    for i, (name, hx) in enumerate(swatches):
        set_cell_shading(table.rows[0].cells[i], hx)
        p = remove_table_cell_paragraph(table.rows[0].cells[i])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        set_run_font(p.add_run(" "), size=8)
        table.rows[0].cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p2 = remove_table_cell_paragraph(table.rows[1].cells[i])
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(4)
        p2.paragraph_format.space_after = Pt(4)
        set_run_font(p2.add_run(f"{name}\n#{hx}"), size=8.3, color=INK, bold=True)
        set_cell_shading(table.rows[1].cells[i], WHITE)
    widths = [1872] * 5
    apply_table_geometry(table, widths, table_width_dxa=CONTENT_W_DXA, indent_dxa=40, cell_margins_dxa={"top": 40, "bottom": 40, "start": 40, "end": 40})
    set_table_borders(table, color=LINE, size=4)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)

    doc.add_heading("3.4 Dirección de audio", level=2)
    add_body(doc, "El audio es el arma principal: instalación, presencia/cuerpo, memoria doméstica y tiempo/repetición forman familias con variantes de distancia, lado, altura, duración, material e intensidad. La música articula recuerdo y presión sin volverse alarma constante. Los eventos fuertes se construyen con dirección, ataque, contraste y recuperación; nunca exigen volumen alto.")


def moodboard_page(doc):
    add_kicker(doc, "VISUAL SYSTEM / REFERENCE BOARD")
    doc.add_heading("3.5 Tablero de dirección visual", level=2)
    add_small(doc, "Referencia conceptual v0.2. No representa arte implementado ni una escena final; fija progresión, materiales, composición y contraste.", italic=True)

    board = ROOT / "Docs/Design/Media/GDD_v0.2_VisualDirectionBoard.jpg"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run()
    shape = run.add_picture(str(board), width=Inches(6.5))
    shape._inline.docPr.set("title", "Echo Protocol - tablero de dirección visual")
    shape._inline.docPr.set(
        "descr",
        "Seis paneles que muestran una instalación clínica invadida gradualmente "
        "por elementos domésticos y un Echo humano cian.",
    )

    add_table(
        doc,
        ["01-02 / Normalidad", "03-04 / Intrusión", "05-06 / Convergencia"],
        [[
            "Clínica legible + un detalle doméstico apagado.",
            "Espacio vacío, dirección fuera de campo y hogar infiltrado.",
            "Echo humano exacto + facility y memoria inseparables.",
        ]],
        [3120, 3120, 3120],
        header_fill=CYAN_DARK,
        zebra=False,
        font_size=8.3,
        compact=True,
    )
    add_small(doc, "Imagen conceptual generada para Echo Protocol el 11/08/2026; uso interno de diseño. Sin logos, texto ni referencias de franquicia.", italic=True)


def technical_mvp_page(doc, bullet_id):
    add_part_opening(doc, "04", "Cronograma de desarrollo", "Tecnología, alcance, hitos y criterios de salida.", page_break_before=True)
    doc.add_heading("4.1 Especificaciones técnicas", level=2)
    add_label_value_table(
        doc,
        [
            ("Engine", "Unity 6.3 LTS"),
            ("Render pipeline", "Universal Render Pipeline (URP)"),
            ("Lenguaje", "C#"),
            ("Input", "Unity Input System / Keyboard + Mouse"),
            ("Perspectiva", "3D / First Person"),
            ("Plataformas objetivo", "Windows y macOS"),
            ("Audio", "Stereo espacial; Headphones / Speakers + Reduced Dynamics planificados"),
            ("Guardado", "Autosave entre habitaciones / capítulos"),
            ("Rendimiento", "Target específico TBD después del vertical slice y profiling"),
        ],
        label_w=2400,
        value_w=6960,
        font_size=9.1,
    )
    doc.add_heading("4.2 MVP", level=2)
    add_table(
        doc,
        ["Incluido", "No MVP"],
        [
            ("Movimiento, cámara e interacción", "Combate, armas o enemigos sistémicos"),
            ("RECALL y reset estable", "Echoes manipulando objetos móviles"),
            ("1 Echo validado; ampliar sólo si un puzzle concreto lo exige", "Inventario complejo o crafting"),
            ("Roster pequeño de pruebas; cantidad final TBD por función de horror", "Multiplayer, networking u online backend"),
            ("Hub pequeño, UI mínima, autosave", "Open world, RPG, achievements o múltiples slots"),
            ("Principio, escalada, clímax y final C", "Localización/ports adicionales no esenciales"),
        ],
        [4680, 4680],
        header_fill=CYAN_DARK,
        font_size=8.75,
    )
    add_callout(doc, "Criterio de scope", "La duración combina puzzle, exploración, audio, narrativa y recuperación. No se agregan pruebas de dos minutos para rellenar 30-45 minutos; cada una debe crear vulnerabilidad y un beat de horror.", fill=MEMORY_LIGHT, accent=MEMORY)


def roadmap_page(doc, bullet_id):
    add_kicker(doc, "04 / Cronograma de desarrollo")
    doc.add_heading("4.3 Estado del prototipo", level=2)
    add_body(doc, "La prueba mecánica original ya fue superada. M0-M5.2 validaron movimiento, interacción, RECALL, Echo exacto y puzzles de una placa, simultaneidad y secuencia. M6.1-M6.3, AD2 y AV1 añadieron una base audiovisual sin cambiar las reglas.")
    add_table(
        doc,
        ["Gate de aceptación", "Resultado esperado"],
        [
            ("Player", "Movimiento, mouse look e interacción human-approved."),
            ("Temporal", "Registro, reset, un Echo y ciclos consecutivos human-approved."),
            ("Puzzle", "M4, M5.1 y M5.2 completables y human-approved."),
            ("Audio", "Causalidad e integración AD2 aprobadas como prototipo; timbres abiertos."),
            ("Visual", "AV1 aprobado como fundamento clínico de prototipo."),
            ("Horror", "Silla/cue entregados técnicamente, pero sin terror validado."),
        ],
        [2900, 6460],
        header_fill=CHARCOAL,
        font_size=8.9,
    )

    doc.add_heading("4.4 Roadmap basado en evidencia", level=2)
    add_table(
        doc,
        ["#", "Hito", "Tiempo", "Entrega propuesta", "Nota"],
        [
            ("M0-M4", "Fundación + loop", "Completo", "07/08/2026", "Core temporal y primer puzzle validados."),
            ("M5", "Puzzle language", "Completo", "09/08/2026", "Simultaneidad y secuencia validadas."),
            ("M6", "Discrepancy proofs", "Parcial", "09/08/2026", "Delivery sí; terror de silla no validado."),
            ("AD2", "Audio prototype", "Completo", "09/08/2026", "Ruta aprobada; timbres reemplazables."),
            ("AV1", "Visual foundation", "Completo", "10/08/2026", "Lugar clínico aprobado como prototipo."),
            ("SD1", "Directional horror", "Actual", "TBD", "Estudio de 5 min; sin puzzle nuevo."),
            ("HM1", "Horror map", "Siguiente", "TBD", "Mapa completo de 30-45 min y vulnerabilidades."),
            ("VS2", "Experience slice", "Planificado", "TBD", "8-10 min; puzzle no basado en placas."),
            ("M7-M8", "Content + polish", "Planificado", "TBD", "Capítulos, QA y builds Win/macOS."),
        ],
        [600, 2100, 1250, 1650, 3760],
        header_fill=CYAN_DARK,
        font_size=7.5,
    )
    add_small(doc, "El cronograma de 16 semanas de la v0.1 quedó superado por el avance real. Las nuevas fechas permanecen TBD; los gates de evidencia y su orden sí son vinculantes.", italic=True)


def risks_tbd_sources_page(doc):
    add_kicker(doc, "04 / Cronograma de desarrollo")
    doc.add_heading("4.5 Riesgos y mitigaciones", level=2)
    add_table(
        doc,
        ["Riesgo", "Impacto", "Mitigación", "Prioridad"],
        [
            ("Reset temporal inconsistente", "Bloquea puzzles y confianza", "Baseline explícito; pruebas de ciclos consecutivos.", "Alta"),
            ("Scope excesivo para una persona", "Juego incompleto", "Roster TBD; recorte antes que pérdida de calidad.", "Alta"),
            ("Anomalías demasiado tempranas", "La regla nunca se vuelve segura", "Gate de puzzles normales antes de M6.", "Alta"),
            ("Puzzles rellenan duración", "Ritmo y tensión caen", "Definir vulnerabilidad y beat antes de producir cada puzzle.", "Alta"),
            ("Audio fuerte o mal localizado", "Molestia física / lectura de bug", "Volumen cómodo; headphones, speakers, mono y Reduced Dynamics.", "Alta"),
            ("Símbolo sin efecto perceptual", "El horror se lee como utilería", "Gate humano: conducta, fuente dudosa, tensión y consecuencia.", "Media"),
            ("Ambigüedad ilegible", "El jugador no entiende nada", "Objetivo 80 % de comprensión; playtest narrativo.", "Media"),
            ("Arte/audio y voz consumen producción", "Menos contenido jugable", "Semi-realismo modular; pocas líneas; assets y voces con licencia.", "Media"),
        ],
        [2450, 2100, 3560, 1250],
        header_fill=CHARCOAL,
        font_size=7.85,
    )

    doc.add_heading("4.6 Decisiones abiertas", level=2)
    add_table(
        doc,
        ["Decisión", "Estado v0.2", "Próximo criterio"],
        [
            ("Nombre del protagonista", "TBD; puede permanecer anónimo", "Sólo definir si mejora narrativa o voz."),
            ("Susurro SD1", "DEFINIDO: 'Can you hear me?' / '¿Me escuchás?'", "Validar dirección, intimidad y lectura trans-pantalla."),
            ("Identidad exacta de la voz", "TBD; puede ser alguien de su vida", "Debe conectar con los traumas sin exposición artificial."),
            ("Detalles concretos de ambos accidentes", "TBD", "Conservar patrón D > A > C y relación hermano/a menor - hijo."),
            ("Sprint", "EN PRUEBA; fuera del inicio", "Evaluar pacing y peligro tardío."),
            ("Echoes con objetos móviles", "NO MVP", "Sólo si un puzzle validado justifica la complejidad."),
            ("Target de rendimiento", "TBD", "Fijar tras vertical slice y profiling en hardware objetivo."),
            ("Nombre / realidad de la instalación", "TBD; institucional y ambigua", "Definir sólo si refuerza la lectura sin explicar la pesadilla."),
            ("Fidelidad del Echo", "DEFINIDO: absoluta para datos capturados", "Cada nuevo canal requiere implementación y tests explícitos."),
            ("Consecuencia visual SD1", "DEFINIDO: sutil; forma exacta TBD", "Probar marca/seam sin collider ni rol de puzzle."),
            ("Mapa de 30-45 min", "TBD", "Fijar capítulos, eventos, recuperación y puzzles por vulnerabilidad."),
            ("Render espacial", "TBD", "Evaluar world/head-locked, parlantes, mono y Reduced Dynamics."),
        ],
        [3100, 2300, 3960],
        header_fill=CYAN_DARK,
        font_size=7.55,
    )

    doc.add_heading("4.7 Fuentes de estructura", level=2)
    p = doc.add_paragraph(style="EP Small")
    set_run_font(p.add_run("Unity Learn - Fill out a game design document: "), size=8.4, color=STEEL)
    add_hyperlink(p, "learn.unity.com/tutorial/fill-out-a-game-design-document", "https://learn.unity.com/tutorial/fill-out-a-game-design-document", size=8.4)
    p = doc.add_paragraph(style="EP Small")
    set_run_font(p.add_run("Unity Learn - Game Design: "), size=8.4, color=STEEL)
    add_hyperlink(p, "learn.unity.com/pathway/game-development/unit/planning-a-game/tutorial/game-design", "https://learn.unity.com/pathway/game-development/unit/planning-a-game/tutorial/game-design", size=8.4)
    p = doc.add_paragraph(style="EP Small")
    set_run_font(p.add_run("Unity - Game Design Document Template PDF: "), size=8.4, color=STEEL)
    add_hyperlink(p, "connect-prd-cdn.unity.com/.../Game-Design-Document-Template.pdf", "https://connect-prd-cdn.unity.com/20201215/83f3733d-3146-42de-8a69-f461d6662eb1/Game-Design-Document-Template.pdf", size=8.4)

    add_callout(doc, "Cierre de versión", "La versión 0.2 sincroniza el GDD con el prototipo y fija SD1 como próximo gate. La siguiente revisión corresponde después del estudio direccional y antes de producir nuevos puzzles o capítulos.", fill=ICE, accent=CYAN_DARK, trailing_space=False)


def configure_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    # Make page 2 display as page 1 while the cover remains unnumbered.
    sectpr = section._sectPr
    pgnum = sectpr.find(qn("w:pgNumType"))
    if pgnum is None:
        pgnum = OxmlElement("w:pgNumType")
        sectpr.append(pgnum)
    pgnum.set(qn("w:start"), "0")

    nil_edges = {
        "top": {"val": "nil", "sz": 0, "color": WHITE},
        "bottom": {"val": "nil", "sz": 0, "color": WHITE},
        "left": {"val": "nil", "sz": 0, "color": WHITE},
        "right": {"val": "nil", "sz": 0, "color": WHITE},
    }

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.line_spacing = Pt(1)
    ht = header.add_table(rows=1, cols=2, width=Inches(6.5))
    apply_table_geometry(ht, [4680, 4680], table_width_dxa=CONTENT_W_DXA, indent_dxa=0, cell_margins_dxa={"top": 0, "bottom": 0, "start": 0, "end": 0})
    for cell in ht.rows[0].cells:
        set_cell_border(cell, **nil_edges)
    lp = remove_table_cell_paragraph(ht.cell(0, 0))
    lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lp.paragraph_format.space_after = Pt(0)
    add_picture_with_alt(
        lp.add_run(),
        BRAND_MARK,
        width=Inches(0.18),
        title="Echo Protocol mark",
        description="Marca de Echo Protocol.",
    )
    set_run_font(lp.add_run("  ECHO PROTOCOL"), name=FONT_DISPLAY, size=8.8, color=CHARCOAL, bold=True, caps=True)
    rp = remove_table_cell_paragraph(ht.cell(0, 1))
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_after = Pt(0)
    set_run_font(rp.add_run("DESIGN DOSSIER  /  GDD 0.2"), name=FONT_MONO, size=7.4, color=STEEL, bold=True)
    for cell in ht.rows[0].cells:
        set_cell_border(cell, bottom={"color": CYAN, "sz": 7}, top={"val": "nil", "sz": 0, "color": WHITE}, left={"val": "nil", "sz": 0, "color": WHITE}, right={"val": "nil", "sz": 0, "color": WHITE})

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fp.paragraph_format.line_spacing = Pt(1)
    ft = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    apply_table_geometry(ft, [6500, 2860], table_width_dxa=CONTENT_W_DXA, indent_dxa=0, cell_margins_dxa={"top": 0, "bottom": 0, "start": 0, "end": 0})
    for cell in ft.rows[0].cells:
        set_cell_border(cell, **nil_edges)
    lp = remove_table_cell_paragraph(ft.cell(0, 0))
    lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lp.paragraph_format.space_after = Pt(0)
    set_run_font(lp.add_run("EP-GDD-001  /  UNIVERSIDAD AUSTRAL"), name=FONT_MONO, size=7.1, color=STEEL)
    rp = remove_table_cell_paragraph(ft.cell(0, 1))
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_after = Pt(0)
    set_run_font(rp.add_run("PÁG. "), name=FONT_MONO, size=7.1, color=STEEL, bold=True)
    add_page_field(rp)

    # Keep first-page header/footer empty.
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""


def set_update_fields(doc: Document):
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    bullet_id, number_id = add_custom_numbering(doc)
    set_update_fields(doc)

    doc.core_properties.title = "Echo Protocol - Game Design Document v0.2"
    doc.core_properties.subject = "Game Design Document"
    doc.core_properties.author = "Lautaro Reinoso"
    doc.core_properties.keywords = "Echo Protocol, GDD, Unity, terror psicológico, temporal puzzle"
    doc.core_properties.comments = "Universidad Austral - Laboratorio de Desarrollo de Videojuegos - Segundo cuatrimestre 2026"

    cover_page(doc)
    add_page_break(doc)
    document_control_page(doc, bullet_id)
    add_page_break(doc)
    intro_page_one(doc)
    add_page_break(doc)
    intro_page_two(doc, bullet_id)
    add_page_break(doc)
    intro_page_three(doc)
    gameplay_overview_page(doc, bullet_id)
    add_page_break(doc)
    player_page(doc, bullet_id)
    recall_page(doc, number_id)
    add_page_break(doc)
    echo_page(doc, bullet_id)
    add_page_break(doc)
    puzzle_page(doc, bullet_id)
    add_page_break(doc)
    flow_level_ui_page(doc, bullet_id)
    add_page_break(doc)
    horror_chapters_page(doc)
    add_page_break(doc)
    narrative_page(doc, bullet_id)
    art_page(doc, bullet_id)
    add_page_break(doc)
    moodboard_page(doc)
    technical_mvp_page(doc, bullet_id)
    roadmap_page(doc, bullet_id)
    add_page_break(doc)
    risks_tbd_sources_page(doc)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
