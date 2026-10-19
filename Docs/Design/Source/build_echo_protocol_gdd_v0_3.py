#!/usr/bin/env python3
"""Build the condensed editable Echo Protocol GDD v0.3 DOCX.

The v0.2 builder remains the visual authority. This version reuses its page,
typography, table, callout, header/footer, brand and moodboard components while
condensing product content to twelve pages.
"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

import build_echo_protocol_gdd_v0_2 as v2


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "Docs/Design/Echo Protocol - Game Design Document v0.3.docx"


def replace_text_in_part(part, replacements):
    for node in part._element.iter(qn("w:t")):
        if not node.text:
            continue
        for old, new in replacements:
            node.text = node.text.replace(old, new)


def replace_version_chrome(doc):
    replacements = [
        ("GAME DESIGN DOCUMENT  0.2", "GAME DESIGN DOCUMENT  0.3"),
        ("VERSIÓN 0.2", "VERSIÓN 0.3"),
        ("11.08.2026", "12.08.2026"),
        ("DESIGN DOSSIER  /  GDD 0.2", "DESIGN DOSSIER  /  GDD 0.3"),
    ]
    replace_text_in_part(doc.part, replacements)
    for section in doc.sections:
        replace_text_in_part(section.header.part, replacements)
        replace_text_in_part(section.footer.part, replacements)


def remove_last_paragraph(doc):
    """Remove a helper-generated caption when it would repeat information."""
    paragraph = doc.paragraphs[-1]
    element = paragraph._element
    element.getparent().remove(element)


def executive_overview(doc, bullet_id):
    v2.add_part_opening(
        doc,
        "00",
        "Lectura ejecutiva",
        "La promesa, el alcance y las reglas que no pueden romperse.",
    )
    v2.add_label_value_table(
        doc,
        [
            ("Proyecto", "Echo Protocol / GDD 0.3 / diseño + prototipo"),
            ("Autor / equipo", "Lautaro Reinoso / 1 desarrollador"),
            ("Género", "Terror psicológico + puzzle temporal + aventura narrativa"),
            ("Formato", "Single-player / 3D / primera persona / Windows y macOS"),
            ("Duración", "30-45 minutos; recortable antes que sacrificar calidad"),
            ("Estado", "Base jugable prototipada; SD1 es la prueba actual de horror espacial"),
        ],
        label_w=2300,
        value_w=7060,
        font_size=9.1,
    )
    doc.add_heading("High concept", level=2)
    v2.add_body(
        doc,
        "El protagonista despierta en una instalación desconocida y extrañamente familiar. "
        "RECALL transforma cada intento en un Echo que repite fielmente sus acciones y coopera "
        "con él. Cuando aparecen ciclos verdaderos que no recuerda, el sistema deja de ser sólo "
        "una herramienta y se convierte en testigo de su negación."
    )
    v2.add_callout(
        doc,
        "North star",
        "Era el juego donde tus versiones pasadas recordaban correctamente lo que vos no podías recordar.",
        fill=v2.CHARCOAL,
        accent=v2.CYAN,
        text_color=v2.WHITE,
    )
    doc.add_heading("Público y tono", level=2)
    v2.add_body(
        doc,
        "Experiencia breve para adolescentes mayores y adultos interesados en terror psicológico, "
        "misterio narrativo y puzzles accesibles. Trabaja muerte, culpa, infancia, accidentes, "
        "pérdida, duelo y memoria distorsionada principalmente por implicación, no por combate o gore."
    )
    v2.add_small(
        doc,
        "Convenciones: VALIDADO = evidencia objetiva o humana; DEFINIDO = dirección aprobada; "
        "EN PRUEBA = hipótesis; TBD = decisión abierta.",
        italic=True,
    )


def pillars_and_promise(doc):
    v2.add_kicker(doc, "01 / Visión de producto")
    doc.add_heading("1.1 Pilares de diseño", level=2)
    v2.add_table(
        doc,
        ["Pilar", "Directriz"],
        [
            ("Trust Before Fear", "La regla se vuelve seguridad antes de que el contexto haga dudar."),
            ("The Mechanic Is the Horror", "RECALL y el Echo exacto exponen lo que el protagonista niega."),
            ("Player-facing Horror", "El miedo debe modificar la percepción o conducta del jugador real."),
            ("Sound as the Main Weapon", "Dirección, proximidad, música, contraste y repetición sostienen el terror."),
            ("Puzzles Create Vulnerability", "Cada prueba obliga a esperar, escuchar, mirar fuera, regresar o comprometerse."),
            ("Small Scope, High Polish", "Un juego corto y coherente vale más que contenido incompleto."),
        ],
        [3000, 6360],
        header_fill=v2.CHARCOAL,
        font_size=8.7,
    )
    doc.add_heading("1.2 Experiencia buscada", level=2)
    v2.add_flow_strip(doc, ["Curiosidad", "Inquietud", "Duda", "Paranoia", "Terror", "Aceptación"])
    v2.add_table(
        doc,
        ["Momento", "Pensamiento buscado"],
        [
            ("Inicio", "Qué interesante esta mecánica."),
            ("Duda espacial", "¿Eso vino del juego o de mi casa?"),
            ("Contradicción", "El Echo hizo exactamente eso. ¿Por qué no lo recuerdo?"),
            ("Revelación", "Si el Echo tiene razón, ¿qué pasó realmente?"),
        ],
        [2200, 7160],
        header_fill=v2.CYAN_DARK,
        font_size=9.0,
    )
    v2.add_callout(
        doc,
        "Alcance",
        "Objetivo: 30-45 minutos muy pulidos. Si hace falta recortar, 20 minutos excelentes tienen prioridad sobre 45 mediocres.",
        fill=v2.MEMORY_LIGHT,
        accent=v2.MEMORY,
    )


def core_loop_and_recall(doc, number_id):
    v2.add_part_opening(doc, "02", "Gameplay", "Cooperación temporal, exposición y verificación.")
    doc.add_heading("2.1 Objetivo y core loop", level=2)
    v2.add_body(
        doc,
        "El objetivo local es resolver pruebas coordinando acciones con un Echo. El objetivo global "
        "es reconstruir dos traumas conectados y confrontar el recuerdo que mantiene activa la pesadilla."
    )
    v2.add_flow_strip(doc, ["Explorar", "Escuchar", "Planificar", "Exponerse", "RECALL", "Cooperar", "Verificar"])
    doc.add_heading("2.2 RECALL", level=2)
    v2.add_body(
        doc,
        "RECALL restaura la prueba a una base conocida y convierte el intento anterior en la única "
        "grabación activa. La espera accidental antes de comenzar no forma parte de la línea temporal; "
        "las pausas deliberadas posteriores sí se reproducen completas."
    )
    v2.add_flow_strip(doc, ["Registrar", "RECALL", "Restaurar", "Crear Echo", "Reproducir", "Continuar"])
    remove_last_paragraph(doc)
    doc.add_heading("Contrato funcional", level=3)
    for item in [
        "Restaurar jugador, mecanismos y prueba a una base conocida.",
        "Reproducir el último intento con timing estable y repetible.",
        "Mantener exactamente un Echo activo.",
        "Soportar resets consecutivos sin acumular estado o degradar la solución.",
    ]:
        v2.add_number(doc, item, number_id)
    v2.add_callout(
        doc,
        "Regla de diseño",
        "El avance combina puzzle, exploración, audio, narrativa, consecuencia y recuperación; la duración no se mide por cantidad de salas.",
        fill=v2.FROST,
        accent=v2.MEMORY,
    )


def echo_contract(doc, bullet_id):
    v2.add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.3 Contrato del Echo", level=2)
    v2.add_callout(
        doc,
        "Regla absoluta",
        "El Echo nunca fabrica ni altera datos que realmente registró. La aparente falta de fiabilidad pertenece a la memoria, el contexto y la percepción.",
        fill=v2.CHARCOAL,
        accent=v2.CYAN,
        text_color=v2.WHITE,
    )
    v2.add_table(
        doc,
        ["Regla confiable", "Horror compatible", "Lectura"],
        [
            ("Repite ruta y timing", "El entorno cambia el significado de la misma ruta", "La memoria queda en duda"),
            ("Termina y desaparece", "Pasos ajenos continúan después", "Presencia no confirmada"),
            ("Reproduce canales implementados", "Un canal nuevo revela más verdad tras ser aprobado", "Testigo más completo"),
            ("Existe si hubo grabación", "Se recupera un ciclo omitido claramente distinto", "Vacío de memoria"),
            ("No actúa por cuenta propia", "La evidencia aparece donde ya estaba orientado", "Relectura inquietante"),
            ("El sonido funcional no miente", "Un motivo temporal responde desde otro espacio", "Contradicción contextual"),
        ],
        [2900, 4180, 2280],
        header_fill=v2.CHARCOAL,
        font_size=8.35,
    )
    doc.add_heading("Capacidades del alcance actual", level=2)
    for item in [
        "Movimiento y orientación del protagonista.",
        "Ocupación física de sensores por Player y Echo.",
        "Sin deformación monstruosa del Echo; el comportamiento y la evidencia llevan el horror.",
    ]:
        v2.add_bullet(doc, item, bullet_id)
    v2.add_body(
        doc,
        "Un ciclo omitido puede revelar una verdad que el protagonista reprimió. Debe presentarse como "
        "otra grabación auténtica, nunca como una copia conocida que cambió arbitrariamente.",
        italic=True,
    )


def puzzles_flow_ui(doc, bullet_id):
    v2.add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.4 Puzzles como vulnerabilidad", level=2)
    v2.add_body(
        doc,
        "La dificultad objetivo es moderada: comprender, planificar y coordinar sin ejecución frame-perfect. "
        "Cada prueba combina una función mecánica con una exposición sensorial concreta."
    )
    v2.add_table(
        doc,
        ["Función", "Situación", "Vulnerabilidad"],
        [
            ("Confianza", "Regla temporal legible", "Aprender que el Echo es exacto"),
            ("Atención dividida", "Terminal o acción sostenida", "Dar la espalda y escuchar"),
            ("Retorno", "Cruzar un espacio conocido", "Comparar lo que cambió"),
            ("Separación", "Echo fuera de vista", "Depender sin confirmar visualmente"),
            ("Escucha", "Esperar mientras algo se aproxima", "Sostener el plan bajo presión"),
            ("Memoria", "Grabación exacta con nuevo contexto", "Aceptar evidencia incómoda"),
        ],
        [1900, 3600, 3860],
        header_fill=v2.CYAN_DARK,
        font_size=8.6,
    )
    doc.add_heading("2.5 Flujo del espacio", level=2)
    v2.add_flow_strip(doc, ["Inicio", "Hub", "Prueba", "RECALL", "Resolución", "Retorno", "Capítulo"])
    v2.add_body(
        doc,
        "Un hub pequeño conecta secuencias lineales y backtracking limitado. Cada retorno puede alterar "
        "un detalle conocido. No hay HP ni muerte convencional; el sueño puede forzar un nuevo ciclo."
    )
    doc.add_heading("2.6 UI, idioma y accesibilidad", level=2)
    for item in [
        "HUD mínimo y preferentemente diegético; sin vida, minimap ni quest tracker.",
        "Voz principal en inglés con subtítulos configurables en español e inglés.",
        "Feedback inequívoco para interacción y RECALL; retículo contextual sólo si mejora legibilidad.",
        "Headphones y Speakers, calibración L/R, controles Ambience/SFX/Voice y Reduced Dynamics planificados.",
        "Mensaje: Designed for headphones. Set a comfortable volume; details do not require high volume.",
    ]:
        v2.add_bullet(doc, item, bullet_id)


def horror_progression(doc):
    v2.add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.7 Progresión de horror", level=2)
    v2.add_body(
        doc,
        "La experiencia establece seguridad mecánica y luego presiona la percepción mediante sonido "
        "espacial, recurrencia, variantes autorales, consecuencias y evidencia incompleta. Puede ser "
        "intensa durante la mayor parte del juego sin confirmar una amenaza física ni corromper el Echo."
    )
    v2.add_table(
        doc,
        ["Capítulo", "Emoción", "Gameplay / horror", "Espacio"],
        [
            ("I - Protocol", "Curiosidad", "Reglas confiables y primer impacto inquietante", "Instalación clínica"),
            ("II - Familiarity", "Inquietud", "Duda espacial y memoria doméstica", "Detalles familiares"),
            ("III - Recollection", "Duda", "Presencia y tiempo invaden las pruebas", "Pasillos personales"),
            ("IV - Denial", "Paranoia", "Grabaciones verdaderas contradicen el relato", "Memoria coexistente"),
            ("V - Acceptance", "Terror / tristeza", "Traumas, responsabilidad y desenlace", "Arquitectura imposible"),
        ],
        [1650, 1450, 3900, 2360],
        header_fill=v2.CYAN_DARK,
        font_size=8.15,
    )
    doc.add_heading("Gramática del miedo", level=2)
    v2.add_table(
        doc,
        ["Priorizar", "Evitar"],
        [
            ("Motivos recurrentes con variantes y consecuencias", "Ruido constante o eventos aleatorios"),
            ("Dirección, proximidad, ataque, espectro y contraste", "Volumen alto como sustituto de diseño"),
            ("Oír, localizar, comprobar y encontrar evidencia parcial", "Anomalías sin preparación ni efecto perceptual"),
            ("Preparación > impacto > consecuencia > recuperación", "Explicación inmediata o persecución constante"),
        ],
        [4680, 4680],
        header_fill=v2.CHARCOAL,
        font_size=8.8,
    )
    v2.add_callout(
        doc,
        "Gate humano",
        "Un evento no funciona porque se reprodujo: debe provocar duda, orientación, chequeo del entorno o una respuesta emocional observable sin exigir volumen inseguro.",
        fill=v2.ALERT_LIGHT,
        accent=v2.ALERT,
    )


def narrative_spine(doc, bullet_id):
    v2.add_kicker(doc, "02 / Gameplay")
    doc.add_heading("2.8 Diseño narrativo", level=2)
    v2.add_callout(
        doc,
        "Giro central",
        "Los Echoes recuerdan correctamente. La memoria y percepción del protagonista son las que han sido alteradas.",
        fill=v2.CHARCOAL,
        accent=v2.CYAN,
        text_color=v2.WHITE,
    )
    v2.add_body(
        doc,
        "La pesadilla conecta dos acontecimientos. Durante la infancia, el protagonista estuvo involucrado "
        "en un accidente que terminó con la muerte de su hermano menor. En la adultez, otro accidente "
        "relacionado con su hijo recreó simbólicamente la misma decisión: provocar indirectamente, entrar "
        "en pánico, huir y alterar el recuerdo. Exactamente qué ocurrió y cuánto podía evitar permanece TBD."
    )
    v2.add_flow_strip(doc, ["Provoca", "Pánico", "Huye", "Altera el recuerdo"])
    remove_last_paragraph(doc)
    doc.add_heading("Voz, figuras y motivos", level=2)
    for item in [
        "Una voz inicialmente profesional guía las pruebas y puede negar hechos observados.",
        "SD1 prueba Can you hear me? / ¿Me escuchás? con una voz juvenil masculina de prototipo; su identidad narrativa definitiva sigue EN PRUEBA.",
        "Figuras del pasado progresan de sombra a silueta, vidrio, reflejo, grabación y reconocimiento parcial.",
        "Cuento, dibujo, melodía, objeto, frase y habitación reaparecen con significado creciente.",
        "No hay NPCs que expliquen la historia de forma convencional.",
    ]:
        v2.add_bullet(doc, item, bullet_id)
    doc.add_heading("Final C - No Awakening Shown", level=2)
    v2.add_body(
        doc,
        "El protagonista confronta el recuerdo, termina la pesadilla y la pantalla corta a negro. No se "
        "muestra si despierta. El cierre busca terror y tristeza sin sumar otro giro que invalide lo anterior."
    )
    v2.add_small(
        doc,
        "La instalación funciona como gramática de la pesadilla. Su nombre y posible contraparte real permanecen TBD.",
        italic=True,
    )


def art_audio_direction(doc, bullet_id):
    v2.add_part_opening(doc, "03", "Dirección audiovisual", "Una clínica que se vuelve íntima, imposible y perturbadora.")
    doc.add_heading("3.1 Evolución visual", level=2)
    v2.add_table(
        doc,
        ["Fase", "Color / material", "Señal narrativa"],
        [
            ("Protocol", "Blanco, gris, azul frío, vidrio y metal", "Control institucional"),
            ("Familiarity", "Calidez pequeña y material doméstico", "Recuerdo infiltrado"),
            ("Recollection", "Madera, textil, fotografía y luz mixta", "Dos mundos coexistiendo"),
            ("Denial / Acceptance", "Paleta contaminada y arquitectura imposible", "Memoria inseparable del espacio"),
        ],
        [1900, 3600, 3860],
        header_fill=v2.CYAN_DARK,
        font_size=8.7,
    )
    doc.add_heading("3.2 Paleta", level=2)
    v2_table = doc.add_table(rows=2, cols=5)
    swatches = [
        ("Institución", v2.ICE),
        ("Sombra", v2.CHARCOAL),
        ("Echo", v2.CYAN),
        ("Hogar", v2.MEMORY),
        ("Herida", v2.ALERT),
    ]
    for index, (name, color) in enumerate(swatches):
        v2.set_cell_shading(v2_table.rows[0].cells[index], color)
        p = v2.remove_table_cell_paragraph(v2_table.rows[0].cells[index])
        p.alignment = v2.WD_ALIGN_PARAGRAPH.CENTER
        v2.set_run_font(p.add_run(" "), size=8)
        p2 = v2.remove_table_cell_paragraph(v2_table.rows[1].cells[index])
        p2.alignment = v2.WD_ALIGN_PARAGRAPH.CENTER
        v2.set_run_font(p2.add_run(f"{name}\n#{color}"), size=8.3, color=v2.INK, bold=True)
    v2.apply_table_geometry(v2_table, [1872] * 5, table_width_dxa=v2.CONTENT_W_DXA, indent_dxa=40, cell_margins_dxa={"top": 50, "bottom": 50, "start": 40, "end": 40})
    v2.set_table_borders(v2_table, color=v2.LINE, size=4)
    doc.add_paragraph()
    doc.add_heading("3.3 Echo y presencia humana", level=2)
    for item in [
        "Silueta humana completa, reconocible y levemente transparente.",
        "Artefactos temporales contenidos; nunca deformación monstruosa progresiva.",
        "Proporciones, materiales, iluminación y composición antes que fotorealismo AAA.",
    ]:
        v2.add_bullet(doc, item, bullet_id)
    doc.add_heading("3.4 Audio como arma principal", level=2)
    v2.add_body(
        doc,
        "Cuatro familias - instalación/estructura, presencia/cuerpo, memoria doméstica y tiempo/repetición - "
        "generan variantes autorales de lado, distancia, altura, duración, material e intensidad. Las fuentes "
        "world-locked respetan el espacio; los momentos head-locked son raros, íntimos y no llevan información "
        "esencial. La música articula recuerdo y presión sin convertirse en alarma constante."
    )
    v2.add_callout(
        doc,
        "Seguridad",
        "El impacto nace de posición, ataque, espectro, contraste, movimiento y sustracción. Auriculares recomendados siempre a volumen cómodo.",
        fill=v2.FROST,
        accent=v2.CYAN_DARK,
    )


def moodboard_appendix(doc):
    v2.add_kicker(doc, "APÉNDICE / REFERENCE BOARD")
    doc.add_heading("Tablero de dirección visual", level=2)
    v2.add_small(
        doc,
        "Referencia conceptual v0.3. No representa una escena final: fija progresión, materiales, composición y contraste.",
        italic=True,
    )
    board = ROOT / "Docs/Design/Media/GDD_v0.2_VisualDirectionBoard.jpg"
    p = doc.add_paragraph()
    p.alignment = v2.WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = v2.Pt(5)
    shape = p.add_run().add_picture(str(board), width=v2.Inches(6.5))
    shape._inline.docPr.set("title", "Echo Protocol - tablero de dirección visual")
    shape._inline.docPr.set("descr", "Seis paneles muestran una instalación clínica invadida gradualmente por elementos domésticos y un Echo humano cian.")
    v2.add_table(
        doc,
        ["01-02 / Normalidad", "03-04 / Intrusión", "05-06 / Convergencia"],
        [[
            "Clínica legible + detalle doméstico apagado.",
            "Dirección fuera de campo + hogar infiltrado.",
            "Echo exacto + instalación y memoria inseparables.",
        ]],
        [3120, 3120, 3120],
        header_fill=v2.CYAN_DARK,
        zebra=False,
        font_size=8.3,
        compact=True,
    )
    v2.add_small(doc, "Imagen conceptual generada para Echo Protocol el 11/08/2026; uso interno de diseño.", italic=True)


def scope_status_roadmap(doc):
    v2.add_part_opening(doc, "04", "Alcance y entrega MVP", "Qué se entrega, qué existe y qué falta validar.")
    doc.add_heading("4.1 Entregable MVP", level=2)
    v2.add_table(
        doc,
        ["Elemento", "Definición para la materia"],
        [
            ("Experiencia", "Single-player de 30-45 min; se recorta antes que sacrificar calidad."),
            ("Mecánica", "Movimiento, interacción, RECALL, reset estable y exactamente un Echo."),
            ("Estructura", "Hub pequeño, cinco capítulos y pruebas que produzcan vulnerabilidad."),
            ("Horror", "Audio espacial, consecuencias perceptibles y Echo siempre confiable."),
            ("Presentación", "UI mínima, subtítulos, accesibilidad sonora y dirección visual coherente."),
            ("Cierre", "Autosave, Final C y builds verificadas para Windows y macOS."),
        ],
        [2600, 6760],
        header_fill=v2.CYAN_DARK,
        font_size=8.8,
    )
    doc.add_heading("4.2 Estado real", level=2)
    v2.add_table(
        doc,
        ["Área", "Estado"],
        [
            ("Player + temporal", "Movimiento, interacción, RECALL, reset y un Echo aprobados humanamente."),
            ("Puzzles", "Secuencias deterministas aprobadas; las placas son lenguaje de prototipo, no roster final."),
            ("Audio", "Mezcla funcional integrada y aprobada como prototipo; los timbres son reemplazables."),
            ("Visual", "Fundamento clínico de prototipo aprobado."),
            ("Horror", "Silla y cue entregados, pero no generaron terror; no son el beat principal."),
        ],
        [2200, 7160],
        header_fill=v2.CHARCOAL,
        font_size=8.6,
    )
    doc.add_heading("4.3 Secuencia para completar el MVP", level=2)
    v2.add_table(
        doc,
        ["Paso", "Estado", "Propósito"],
        [
            ("SD1", "ACTUAL", "Estudio direccional de 5 min; sin puzzle nuevo ni corrupción del Echo."),
            ("Mapa MVP", "SIGUIENTE", "Recorrido completo, vulnerabilidades, eventos y recuperación."),
            ("Slice evaluable", "PLANIFICADO", "8-10 min con un puzzle no basado en placas y horror integrado."),
            ("Entrega final", "TBD", "Cinco capítulos, QA, accesibilidad y builds Windows/macOS."),
        ],
        [1800, 1700, 5860],
        header_fill=v2.CYAN_DARK,
        font_size=8.65,
    )


def risks_and_open_decisions(doc):
    kicker = v2.add_kicker(doc, "04 / Alcance y entrega MVP")
    doc.add_heading("4.4 Riesgos principales", level=2)
    v2.add_table(
        doc,
        ["Riesgo", "Mitigación"],
        [
            ("Scope excesivo para una persona", "Recortar antes que perder coherencia o acabado."),
            ("Puzzles usados como relleno", "Definir vulnerabilidad y beat de horror antes de producir."),
            ("Horror detectable pero emocionalmente plano", "Gate humano sin priming; iterar timing, motivo y consecuencia."),
            ("Audio fuerte o mal localizado", "Volumen cómodo; headphones, speakers, mono y Reduced Dynamics."),
            ("Ambigüedad narrativa ilegible", "Evidencia causal suficiente y playtests de reconstrucción."),
            ("Arte, voz y assets consumen producción", "Reutilización modular, pocas líneas, costo cero y licencia trazable."),
        ],
        [3300, 6060],
        header_fill=v2.CHARCOAL,
        font_size=8.65,
    )
    doc.add_heading("4.5 Decisiones abiertas de producto", level=2)
    v2.add_table(
        doc,
        ["Decisión", "Estado", "Próximo criterio"],
        [
            ("Identidad final de la voz", "EN PRUEBA", "Conectar al hermano menor sin exposición ni casting artificial."),
            ("Detalles de los dos accidentes", "TBD", "Conservar culpa, patrón de huida y vínculo hermano-hijo."),
            ("Mapa de 30-45 min", "TBD", "Fijar capítulos, eventos, puzzles y recuperación."),
            ("Roster de puzzles", "TBD", "Cada uno debe fabricar una vulnerabilidad distinta."),
            ("Consecuencia visual SD1", "Sutil / TBD", "Debe ser verificable y no cumplir función mecánica."),
            ("Render espacial y accesibilidad", "TBD", "Validar headphones, speakers, mono y Reduced Dynamics."),
        ],
        [2750, 1900, 4710],
        header_fill=v2.CYAN_DARK,
        font_size=8.15,
    )
    v2.add_callout(
        doc,
        "Cierre de versión",
        "La v0.3 concentra la brújula de producto en doce páginas. El detalle implementado, la arquitectura, los playtests y el roadmap vivo permanecen en los documentos del repositorio.",
        fill=v2.ICE,
        accent=v2.CYAN_DARK,
        trailing_space=False,
    )


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    v2.configure_page(doc)
    v2.configure_styles(doc)
    bullet_id, number_id = v2.add_custom_numbering(doc)
    v2.set_update_fields(doc)

    doc.core_properties.title = "Echo Protocol - Game Design Document v0.3"
    doc.core_properties.subject = "Condensed Game Design Document"
    doc.core_properties.author = "Lautaro Reinoso"
    doc.core_properties.keywords = "Echo Protocol, GDD, Unity, terror psicológico, temporal puzzle"
    doc.core_properties.comments = "Universidad Austral - Laboratorio de Desarrollo de Videojuegos - Segundo cuatrimestre 2026"

    v2.cover_page(doc)
    v2.add_page_break(doc)
    executive_overview(doc, bullet_id)
    v2.add_page_break(doc)
    pillars_and_promise(doc)
    v2.add_page_break(doc)
    core_loop_and_recall(doc, number_id)
    v2.add_page_break(doc)
    echo_contract(doc, bullet_id)
    v2.add_page_break(doc)
    puzzles_flow_ui(doc, bullet_id)
    v2.add_page_break(doc)
    horror_progression(doc)
    v2.add_page_break(doc)
    narrative_spine(doc, bullet_id)
    v2.add_page_break(doc)
    art_audio_direction(doc, bullet_id)
    v2.add_page_break(doc)
    moodboard_appendix(doc)
    v2.add_page_break(doc)
    scope_status_roadmap(doc)
    risks_and_open_decisions(doc)

    replace_version_chrome(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
